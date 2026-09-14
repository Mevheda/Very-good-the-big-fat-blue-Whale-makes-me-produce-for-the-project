from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_run_font(run, *, east_asia: str, size: int, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def _set_paragraph_format(
    paragraph,
    *,
    indent: bool = True,
    hanging: bool = False,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if hanging:
        fmt.left_indent = Cm(0.74)
        fmt.first_line_indent = Cm(-0.74)
    elif indent:
        fmt.first_line_indent = Cm(0.74)


def build_docx(
    assignment: dict[str, Any],
    questions: list[dict[str, Any]],
    answers: dict[str, list[dict[str, Any]]],
    output_path: Path,
) -> None:
    missing = [
        str(question["id"])
        for question in questions
        if str(question["id"]) not in answers or not answers[str(question["id"])]
    ]
    if missing:
        raise ValueError(f"Missing answers for questions: {', '.join(missing)}")

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run(assignment.get("document_title", assignment["title"]))
    _set_run_font(title_run, east_asia="黑体", size=16, bold=True)

    subtitle_text = str(assignment.get("subtitle", "")).strip()
    if subtitle_text:
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(10)
        subtitle_run = subtitle.add_run(subtitle_text)
        _set_run_font(subtitle_run, east_asia="宋体", size=12)

    section_title = str(assignment.get("section_title", "")).strip()
    if section_title:
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.space_before = Pt(2)
        heading.paragraph_format.space_after = Pt(8)
        heading_run = heading.add_run(section_title)
        _set_run_font(heading_run, east_asia="黑体", size=14, bold=True)

    for question in questions:
        question_id = str(question["id"])
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(5)
        heading.paragraph_format.space_after = Pt(2)
        heading_run = heading.add_run(f"{question_id}、{question['text']}")
        _set_run_font(heading_run, east_asia="宋体", size=12, bold=True)

        answer_label = document.add_paragraph()
        answer_label.paragraph_format.space_after = Pt(0)
        label_run = answer_label.add_run("答：")
        _set_run_font(label_run, east_asia="黑体", size=12, bold=True)

        numbered = 0
        for item in answers[question_id]:
            text = str(item.get("text", "")).strip()
            style = str(item.get("style", "paragraph"))
            if not text:
                continue
            if style == "bullet":
                numbered += 1
                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"{numbered}、{text}")
                _set_run_font(run, east_asia="宋体", size=12)
                _set_paragraph_format(paragraph, indent=False, hanging=True)
            elif style == "subheading":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(text)
                _set_run_font(run, east_asia="黑体", size=12, bold=True)
                _set_paragraph_format(paragraph, indent=False)
            else:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(text)
                _set_run_font(run, east_asia="宋体", size=12)
                _set_paragraph_format(paragraph, indent=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def build_from_files(
    assignment_json: Path,
    questions_json: Path,
    answers_json: Path,
    output_path: Path,
) -> None:
    assignment = json.loads(assignment_json.read_text(encoding="utf-8"))
    questions = json.loads(questions_json.read_text(encoding="utf-8"))
    answers = json.loads(answers_json.read_text(encoding="utf-8"))
    build_docx(assignment, questions, answers, output_path)
